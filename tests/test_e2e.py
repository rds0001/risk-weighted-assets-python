import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from rwa_engine.pipeline import run_dataset
from rwa_engine.synthetic import generate_synthetic_dataset


def test_full_run_creates_versioned_excel_package(tmp_path: Path):
    dataset = generate_synthetic_dataset(tmp_path, seed=5752026)
    assert dataset.parts[-2:] == ("2026-08-31", "v1.0.0")
    input_files = list((dataset / "inputs").glob("RWA_IN_*.xlsx"))
    assert len(input_files) == 16
    for path in input_files:
        wb = load_workbook(path, read_only=False)
        assert {"README", "DATA_DICTIONARY", "_LOOKUPS", "CHANGELOG"}.issubset(wb.sheetnames)
        assert wb["_LOOKUPS"].sheet_state == "hidden"
        assert all(ws.tables for ws in wb.worksheets)
    output = run_dataset(dataset)
    files = sorted(output.glob("RWA_OUT_*.xlsx"))
    assert len(files) == 6
    manifest = json.loads((output / "run_manifest.json").read_text())
    assert manifest["metrics"]["TREA"] > 0
    assert manifest["metrics"]["CET1"] > 0
    assert manifest["controls_passed"] == manifest["control_count"]
    assert {"input_hash", "code_hash", "calculation_fingerprint", "engine_version"}.issubset(manifest)
    for path in files:
        wb = load_workbook(path, read_only=False)
        assert "Run_Info" in wb.sheetnames
        assert "tbl_out_run_info" in wb["Run_Info"].tables
    summary = pd.read_excel(files[0], sheet_name="Metrics")
    assert {"APPLIED", "FULLY_LOADED"} == set(summary["view"])
    applied = pd.read_excel(files[0], sheet_name="Applied_TREA").iloc[0]
    fully_loaded = pd.read_excel(files[0], sheet_name="Fully_Loaded_TREA").iloc[0]
    assert not bool(applied["floor_binding"])
    assert bool(fully_loaded["floor_binding"])
    audit = pd.read_excel(files[-1], sheet_name="Validation_Issues")
    controls = pd.read_excel(files[-1], sheet_name="Reconciliations")
    assert set(audit["code"]) <= {"LEGACY_SA_SUPPORTING_FACTOR"}
    assert not (audit["severity"] == "ERROR").any()
    assert controls["passed"].all()
    pillar1 = pd.read_excel(files[1], sheet_name="SA_Detail")
    assert {
        "calculation_run_id",
        "result_version",
        "result_as_of_date",
        "result_knowledge_time",
        "result_rule_set_id",
        "result_view",
        "result_is_official",
    }.issubset(pillar1.columns)
    assert pillar1["result_is_official"].all()
