"""Read-only Katalogsicht auf versionierte RWA-Datensätze und Läufe."""

from __future__ import annotations

import json
import re
from datetime import datetime, timezone
from hashlib import sha256
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from rwa_engine.contracts import TABLE_SPECS

_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
_VERSION = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._-]*$")


class CatalogError(ValueError):
    """Ungültiger oder nicht vorhandener Katalogpfad."""


def _json(path: Path) -> dict[str, Any]:
    try:
        value = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise CatalogError(f"Manifest nicht lesbar: {path.name}") from exc
    if not isinstance(value, dict):
        raise CatalogError(f"Manifest ist kein Objekt: {path.name}")
    return value


def _time(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat()


def _size(path: Path) -> int:
    return path.stat().st_size


def _sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _audit_counts(path: Path) -> tuple[int, int]:
    """Liest nur Severity aus dem Audit-Workbook, ohne Ergebnisdaten zu laden."""
    workbook = None
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        if "Validation_Issues" not in workbook.sheetnames:
            return 0, 0
        rows = workbook["Validation_Issues"].iter_rows(values_only=True)
        header = next(rows, ())
        severity_index = list(header).index("severity")
        severities = [
            str(row[severity_index]).upper() for row in rows if row and row[severity_index] is not None
        ]
        return severities.count("ERROR"), severities.count("WARNING")
    except (OSError, ValueError, StopIteration):
        return 0, 0
    finally:
        if workbook is not None:
            workbook.close()


class DataCatalog:
    """Sicherer Zugriff auf ``daten/rechenlaeufe/YYYY-MM-DD/version``."""

    def __init__(self, data_root: Path):
        self.data_root = data_root.resolve()

    def dataset_path(self, identifier: str) -> Path:
        parts = identifier.split("/")
        if len(parts) != 2 or not _DATE.fullmatch(parts[0]) or not _VERSION.fullmatch(parts[1]):
            raise CatalogError("Datensatz-ID muss YYYY-MM-DD/version entsprechen")
        path = (self.data_root / parts[0] / parts[1]).resolve()
        try:
            path.relative_to(self.data_root)
        except ValueError as exc:
            raise CatalogError("Datensatz liegt außerhalb der Datenwurzel") from exc
        if not path.is_dir() or not (path / "inputs").is_dir():
            raise CatalogError("Datensatz nicht gefunden")
        return path

    def datasets(self) -> list[dict[str, Any]]:
        rows: list[dict[str, Any]] = []
        if not self.data_root.is_dir():
            return rows
        for day in self.data_root.iterdir():
            if not day.is_dir() or not _DATE.fullmatch(day.name):
                continue
            for version in day.iterdir():
                identifier = f"{day.name}/{version.name}"
                if (
                    not version.is_dir()
                    or not _VERSION.fullmatch(version.name)
                    or not (version / "inputs").is_dir()
                ):
                    continue
                try:
                    manifest = _json(version / "dataset_manifest.json")
                except CatalogError:
                    manifest = {}
                runs = self._runs(version)
                rows.append(
                    {
                        "id": identifier,
                        "as_of_date": manifest.get("as_of_date", day.name),
                        "version": manifest.get("dataset_version", version.name),
                        "bank_profile": manifest.get("bank_profile", "UNBEKANNT"),
                        "generated_at": manifest.get("generated_at"),
                        "input_count": len(list((version / "inputs").glob("*.xlsx"))),
                        "run_count": len(runs),
                        "latest_run": runs[0] if runs else None,
                    }
                )
        return sorted(rows, key=lambda row: (row["as_of_date"], row["version"]), reverse=True)

    def detail(self, identifier: str) -> dict[str, Any]:
        dataset = self.dataset_path(identifier)
        manifest = (
            _json(dataset / "dataset_manifest.json") if (dataset / "dataset_manifest.json").is_file() else {}
        )
        counts = manifest.get("table_row_counts", {})
        expected_hashes = manifest.get("input_sha256", {})
        workbook_counts: dict[str, int] = {}
        workbook_tables: dict[str, int] = {}
        for logical_name, spec in TABLE_SPECS.items():
            workbook_counts[spec.workbook] = workbook_counts.get(spec.workbook, 0) + int(
                counts.get(logical_name, 0)
            )
            workbook_tables[spec.workbook] = workbook_tables.get(spec.workbook, 0) + 1
        expected = set(manifest.get("input_workbooks", []))
        inputs = []
        for path in sorted((dataset / "inputs").glob("*.xlsx")):
            expected_hash = expected_hashes.get(path.name)
            inputs.append(
                {
                    "name": path.name,
                    "size_bytes": _size(path),
                    "modified_at": _time(path),
                    "row_count": workbook_counts.get(path.name),
                    "table_count": workbook_tables.get(path.name),
                    "declared": not expected or path.name in expected,
                    "integrity": "UNDECLARED"
                    if expected and path.name not in expected
                    else "CHANGED"
                    if expected_hash and _sha256(path) != expected_hash
                    else "MATCH",
                }
            )
        present = {item["name"] for item in inputs}
        runs = self._runs(dataset)
        return {
            "id": identifier,
            "manifest": {
                "as_of_date": manifest.get("as_of_date", identifier.split("/")[0]),
                "dataset_version": manifest.get("dataset_version", identifier.split("/")[1]),
                "bank_profile": manifest.get("bank_profile", "UNBEKANNT"),
                "profile_version": manifest.get("profile_version"),
                "generated_at": manifest.get("generated_at"),
                "table_count": len(counts),
                "row_count": sum(int(value) for value in counts.values()),
                "expected_input_count": len(expected) if expected else len(inputs),
                "missing_inputs": sorted(expected - present),
                "changed_inputs": [item["name"] for item in inputs if item["integrity"] == "CHANGED"],
            },
            "inputs": inputs,
            "runs": runs,
        }

    def _runs(self, dataset: Path) -> list[dict[str, Any]]:
        runs: list[dict[str, Any]] = []
        output_root = dataset / "outputs"
        if not output_root.is_dir():
            return runs
        for run_dir in output_root.iterdir():
            if not run_dir.is_dir() or not run_dir.name.startswith("RUN-"):
                continue
            manifest_path = run_dir / "run_manifest.json"
            if not manifest_path.is_file():
                rejected = sorted(run_dir.glob("*_REJECTED.xlsx"))
                if rejected:
                    validation_errors, validation_warnings = _audit_counts(rejected[-1])
                    runs.append(
                        {
                            "run_id": run_dir.name,
                            "status": "REJECTED",
                            "created_at": _time(rejected[-1]),
                            "controls_passed": 0,
                            "control_count": 0,
                            "metrics": {},
                            "validation_errors": validation_errors,
                            "validation_warnings": validation_warnings,
                            "files": [self._file(rejected[-1])],
                        }
                    )
                continue
            try:
                manifest = _json(manifest_path)
            except CatalogError:
                continue
            declared = manifest.get("output_files", [])
            files = [self._file(run_dir / name) for name in declared if (run_dir / name).is_file()]
            audit = next(
                (
                    run_dir / name
                    for name in declared
                    if name.endswith("_05_Audit.xlsx") and (run_dir / name).is_file()
                ),
                None,
            )
            validation_errors, validation_warnings = _audit_counts(audit) if audit else (0, 0)
            runs.append(
                {
                    "run_id": manifest.get("run_id", run_dir.name),
                    "status": manifest.get("status", "UNBEKANNT"),
                    "created_at": manifest.get("created_at", _time(manifest_path)),
                    "engine_version": manifest.get("engine_version"),
                    "rule_set_id": manifest.get("rule_set_id"),
                    "controls_passed": manifest.get("controls_passed", 0),
                    "control_count": manifest.get("control_count", 0),
                    "validation_errors": validation_errors,
                    "validation_warnings": validation_warnings,
                    "metrics": manifest.get("metrics", {}),
                    "files": files,
                }
            )
        return sorted(runs, key=lambda row: str(row.get("created_at") or ""), reverse=True)

    @staticmethod
    def _file(path: Path) -> dict[str, Any]:
        return {"name": path.name, "size_bytes": _size(path), "modified_at": _time(path)}

    def file_path(self, identifier: str, area: str, filename: str, run_id: str | None = None) -> Path:
        dataset = self.dataset_path(identifier)
        if Path(filename).name != filename or not filename.lower().endswith(".xlsx"):
            raise CatalogError("Ungültiger Excel-Dateiname")
        if area == "input":
            parent = dataset / "inputs"
        elif area == "output" and run_id and Path(run_id).name == run_id and run_id.startswith("RUN-"):
            parent = dataset / "outputs" / run_id
        else:
            raise CatalogError("Ungültiger Dateibereich")
        path = (parent / filename).resolve()
        try:
            path.relative_to(parent.resolve())
        except ValueError as exc:
            raise CatalogError("Datei liegt außerhalb des erlaubten Bereichs") from exc
        if not path.is_file():
            raise CatalogError("Datei nicht gefunden")
        return path
