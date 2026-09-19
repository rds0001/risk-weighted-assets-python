"""Excel-Ein-/Ausgabe mit strukturierten Tabellen und sauberer Formatierung."""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Mapping

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .contracts import COMMON_COLUMNS, TABLE_SPECS, TableSpec, specs_for_workbook
from .supporting import IRB_SUPPORT_FIELDS, SUPPORT_FIELDS, normalise_support_columns, support_issues


@dataclass
class ValidationIssue:
    severity: str
    code: str
    table: str
    row_ref: str
    field: str
    message: str


def _safe_excel_value(value):
    if pd.isna(value):
        return None
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _format_workbook(path: Path, specs: Mapping[str, TableSpec]) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    meta_fill = PatternFill("solid", fgColor="D9EAF7")
    for logical_name, spec in specs.items():
        ws = wb[spec.sheet]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if cell.value in COMMON_COLUMNS:
                cell.fill = meta_fill
                cell.font = Font(color="000000", bold=True)
        max_row = max(ws.max_row, 2)
        max_col = max(ws.max_column, 1)
        ref = f"A1:{ws.cell(max_row, max_col).coordinate}"
        if ws.max_row == 1:
            # Eine leere Datenzeile hält eine gültige strukturierte Tabelle vor.
            for col in range(1, max_col + 1):
                ws.cell(2, col, None)
        tab = Table(displayName=spec.table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tab)
        for col_cells in ws.columns:
            values = [str(c.value) for c in col_cells[: min(len(col_cells), 200)] if c.value is not None]
            width = min(max([len(v) for v in values] + [12]) + 2, 42)
            ws.column_dimensions[col_cells[0].column_letter].width = width
        for cell in ws[1]:
            if cell.value in {"as_of_date", "valid_from", "valid_to", "known_from", "known_to"}:
                for row in range(2, ws.max_row + 1):
                    ws.cell(row, cell.column).number_format = "yyyy-mm-dd"
            if cell.value == "is_official":
                dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=False)
                ws.add_data_validation(dv)
                dv.add(f"{cell.column_letter}2:{cell.column_letter}1048576")
    for sheet_name, table_name in [
        ("README", "tbl_doc_readme"),
        ("DATA_DICTIONARY", "tbl_doc_dictionary"),
        ("_LOOKUPS", "tbl_doc_lookups"),
        ("CHANGELOG", "tbl_doc_changelog"),
    ]:
        ws = wb[sheet_name]
        ws.freeze_panes = "A2"
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
        ref = f"A1:{ws.cell(max(ws.max_row, 2), max(ws.max_column, 1)).coordinate}"
        tab = Table(displayName=table_name, ref=ref)
        tab.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        ws.add_table(tab)
        for col_cells in ws.columns:
            values = [str(c.value) for c in col_cells if c.value is not None]
            ws.column_dimensions[col_cells[0].column_letter].width = min(
                max([len(v) for v in values] + [12]) + 2, 60
            )
    wb["_LOOKUPS"].sheet_state = "hidden"
    wb.save(path)


def write_input_workbooks(directory: Path, tables: Mapping[str, pd.DataFrame]) -> list[Path]:
    directory.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []
    for workbook in sorted({s.workbook for s in TABLE_SPECS.values()}):
        path = directory / workbook
        specs = specs_for_workbook(workbook)
        with pd.ExcelWriter(path, engine="openpyxl", datetime_format="yyyy-mm-dd") as writer:
            pd.DataFrame(
                [
                    {"property": "workbook", "value": workbook},
                    {"property": "purpose", "value": "Kanonische, institutsneutrale RWA-Eingabedaten"},
                    {
                        "property": "editing_rule",
                        "value": "Nur Datenzellen ändern; Spalten- und Tabellennamen stabil halten",
                    },
                    {
                        "property": "versioning",
                        "value": "Korrekturen als neue append-only Version, nicht überschreiben",
                    },
                    {
                        "property": "official_default",
                        "value": "Neue aktive Ist-Datensätze standardmäßig TRUE",
                    },
                ]
            ).to_excel(writer, sheet_name="README", index=False)
            dictionary = []
            for logical_name, spec in specs.items():
                for col in spec.all_columns:
                    dictionary.append(
                        {
                            "logical_table": logical_name,
                            "sheet": spec.sheet,
                            "excel_table": spec.table_name,
                            "field": col,
                            "required": col not in (IRB_SUPPORT_FIELDS if logical_name == "irb_parameter" else SUPPORT_FIELDS if logical_name == "sa_classification" else {}),
                            "description": col.replace("_", " "),
                        }
                    )
            pd.DataFrame(dictionary).to_excel(writer, sheet_name="DATA_DICTIONARY", index=False)
            pd.DataFrame(
                {
                    "domain": [
                        "record_status",
                        "record_status",
                        "record_status",
                        "record_status",
                        "boolean",
                        "boolean",
                    ],
                    "value": ["ACTIVE", "PROVISIONAL", "SIMULATED", "CANCELLED", "TRUE", "FALSE"],
                }
            ).to_excel(writer, sheet_name="_LOOKUPS", index=False)
            pd.DataFrame(
                [
                    {
                        "template_version": "1.2.0",
                        "change_date": date.today(),
                        "change": "Additive optionale Unterstützungsfaktoren für SA und IRB",
                    }
                ]
            ).to_excel(writer, sheet_name="CHANGELOG", index=False)
            for logical_name, spec in specs.items():
                frame = tables.get(logical_name, pd.DataFrame(columns=spec.all_columns)).copy()
                frame = normalise_support_columns(frame, logical_name)
                for col in spec.all_columns:
                    if col not in frame.columns:
                        frame[col] = None
                frame = frame.loc[:, spec.all_columns]
                frame.to_excel(writer, sheet_name=spec.sheet, index=False)
        _format_workbook(path, specs)
        written.append(path)
    return written


def read_input_workbooks(directory: Path) -> tuple[dict[str, pd.DataFrame], list[ValidationIssue]]:
    tables: dict[str, pd.DataFrame] = {}
    issues: list[ValidationIssue] = []
    for logical_name, spec in TABLE_SPECS.items():
        path = directory / spec.workbook
        if not path.exists():
            issues.append(ValidationIssue("ERROR", "MISSING_WORKBOOK", logical_name, "", "", str(path)))
            continue
        try:
            frame = pd.read_excel(path, sheet_name=spec.sheet, engine="openpyxl")
        except ValueError as exc:
            issues.append(ValidationIssue("ERROR", "MISSING_SHEET", logical_name, "", "", str(exc)))
            continue
        frame = normalise_support_columns(frame, logical_name)
        missing = [c for c in spec.all_columns if c not in frame.columns]
        if missing:
            issues.append(
                ValidationIssue(
                    "ERROR", "MISSING_COLUMNS", logical_name, "", ",".join(missing), "Pflichtspalten fehlen"
                )
            )
            continue
        frame = frame.loc[:, spec.all_columns].dropna(how="all")
        for col in ("as_of_date", "valid_from", "valid_to", "known_from", "known_to"):
            frame[col] = pd.to_datetime(frame[col], errors="coerce")
        if "is_official" in frame:
            frame["is_official"] = frame["is_official"].map(
                lambda x: x if isinstance(x, bool) else str(x).strip().upper() in {"TRUE", "1", "YES", "JA"}
            )
        tables[logical_name] = frame
    return tables, issues


def select_official_as_of(frame: pd.DataFrame, as_of_date: date, knowledge_time: datetime) -> pd.DataFrame:
    if frame.empty:
        return frame.copy()
    frame = frame.copy()
    for col in ("as_of_date", "valid_from", "valid_to", "known_from", "known_to"):
        frame[col] = pd.to_datetime(frame[col], errors="coerce")
    as_of = pd.Timestamp(as_of_date)
    known = pd.Timestamp(knowledge_time)
    mask = (
        frame["is_official"].fillna(False)
        & (frame["valid_from"].isna() | (frame["valid_from"] <= as_of))
        & (frame["valid_to"].isna() | (frame["valid_to"] > as_of))
        & (frame["known_from"].isna() | (frame["known_from"] <= known))
        & (frame["known_to"].isna() | (frame["known_to"] > known))
        & ~frame["record_status"].isin(["CANCELLED", "SUPERSEDED"])
    )
    selected = frame.loc[mask].copy()
    if selected.empty:
        return selected
    selected = selected.sort_values(["business_key", "version_no", "known_from"])
    return selected.groupby("business_key", as_index=False, sort=False).tail(1).reset_index(drop=True)


def validate_tables(tables: Mapping[str, pd.DataFrame]) -> list[ValidationIssue]:
    issues: list[ValidationIssue] = []
    for logical_name, spec in TABLE_SPECS.items():
        frame = tables.get(logical_name)
        if frame is None:
            continue
        if spec.required and frame.empty:
            issues.append(
                ValidationIssue(
                    "WARNING",
                    "EMPTY_TABLE",
                    logical_name,
                    "",
                    "",
                    "Tabelle ist leer; Modul kann NOT_APPLICABLE sein",
                )
            )
        if not frame.empty:
            mandatory = {
                "record_id",
                "business_key",
                "version_no",
                "as_of_date",
                "valid_from",
                "known_from",
                "is_official",
                "record_status",
                spec.columns[0],
            }
            for field in sorted(mandatory & set(frame.columns)):
                missing_value = frame[field].isna() | frame[field].astype(str).str.strip().eq("")
                for index in frame.index[missing_value][:20]:
                    issues.append(
                        ValidationIssue(
                            "ERROR",
                            "MISSING_REQUIRED_VALUE",
                            logical_name,
                            str(frame.at[index, "record_id"]),
                            field,
                            "Pflichtwert fehlt",
                        )
                    )
            dup = frame[frame.duplicated(list(spec.primary_key), keep=False)]
            for _, row in dup.head(20).iterrows():
                issues.append(
                    ValidationIssue(
                        "ERROR",
                        "DUPLICATE_KEY",
                        logical_name,
                        str(row.get("record_id", "")),
                        "record_id",
                        "Schlüssel nicht eindeutig",
                    )
                )
            bad_versions = pd.to_numeric(frame["version_no"], errors="coerce").isna()
            if bad_versions.any():
                issues.append(
                    ValidationIssue(
                        "ERROR",
                        "INVALID_VERSION",
                        logical_name,
                        "",
                        "version_no",
                        "Version muss numerisch sein",
                    )
                )
            if "valid_to" in frame:
                invalid_interval = frame["valid_to"].notna() & (
                    pd.to_datetime(frame["valid_to"], errors="coerce")
                    <= pd.to_datetime(frame["valid_from"], errors="coerce")
                )
                for index in frame.index[invalid_interval][:20]:
                    issues.append(
                        ValidationIssue(
                            "ERROR",
                            "INVALID_VALIDITY_INTERVAL",
                            logical_name,
                            str(frame.at[index, "record_id"]),
                            "valid_to",
                            "valid_to muss nach valid_from liegen",
                        )
                    )
    # Zentrale Referenzprüfungen
    refs = [
        ("exposure_lot", "party_id", "party", "party_id"),
        ("exposure_lot", "contract_id", "product_contract", "contract_id"),
        ("sa_classification", "exposure_id", "exposure_lot", "exposure_id"),
        ("irb_parameter", "exposure_id", "exposure_lot", "exposure_id"),
        ("protection_allocation", "exposure_id", "exposure_lot", "exposure_id"),
        ("derivative_trade", "netting_set_id", "netting_set", "netting_set_id"),
        ("securitisation_tranche", "pool_id", "securitisation_pool", "pool_id"),
        ("cashflow", "contract_id", "product_contract", "contract_id"),
    ]
    for child, child_col, parent, parent_col in refs:
        c, p = tables.get(child), tables.get(parent)
        if c is None or p is None or c.empty:
            continue
        missing = set(c[child_col].dropna().astype(str)) - set(p[parent_col].dropna().astype(str))
        for key in sorted(missing)[:50]:
            issues.append(
                ValidationIssue(
                    "ERROR", "BROKEN_REFERENCE", child, key, child_col, f"Kein {parent}.{parent_col}={key}"
                )
            )
    bounded_fields = {
        "pd_estimate": (0.0, 1.0),
        "pd_floor": (0.0, 1.0),
        "lgd_estimate": (0.0, 1.0),
        "lgd_floor": (0.0, 1.0),
        "correlation": (-1.0, 1.0),
        "ownership_share": (0.0, 1.0),
        "haircut": (0.0, 1.0),
        "probability": (0.0, 1.0),
        "weight": (0.0, 1.0),
        "eligibility_factor": (0.0, 1.0),
        "cet1_share": (0.0, 1.0),
        "tier1_share": (0.0, 1.0),
        "intra_bucket_correlation": (-1.0, 1.0),
        "inter_bucket_correlation": (-1.0, 1.0),
        "hedge_correlation": (-1.0, 1.0),
        "npe_share": (0.0, 1.0),
        "attachment": (0.0, 1.0),
        "detachment": (0.0, 1.0),
    }
    for logical_name, frame in tables.items():
        if frame is None or frame.empty:
            continue
        for field, (lower, upper) in bounded_fields.items():
            if field not in frame:
                continue
            numeric = pd.to_numeric(frame[field], errors="coerce")
            invalid = frame[field].notna() & (numeric.isna() | (numeric < lower) | (numeric > upper))
            for index in frame.index[invalid][:20]:
                issues.append(
                    ValidationIssue(
                        "ERROR",
                        "VALUE_OUT_OF_RANGE",
                        logical_name,
                        str(frame.at[index, "record_id"]),
                        field,
                        f"Wert muss im Intervall [{lower}, {upper}] liegen",
                    )
                )
    params = tables.get("regulatory_parameter")
    if params is not None and not params.empty:
        key_columns = ["parameter_key", "dimension_1", "dimension_2", "rule_set_id"]
        normalized = params[key_columns].fillna("").astype(str)
        duplicated = normalized.duplicated(key_columns, keep=False)
        for index in params.index[duplicated][:20]:
            issues.append(
                ValidationIssue(
                    "ERROR",
                    "DUPLICATE_PARAMETER",
                    "regulatory_parameter",
                    str(params.at[index, "record_id"]),
                    "parameter_key",
                    "Fachparameter ist innerhalb des Regelsatzes nicht eindeutig",
                )
            )
    run_config = tables.get("run_config")
    if run_config is not None and not run_config.empty:
        required_keys = {
            "as_of_date",
            "knowledge_time",
            "consolidation_scope_id",
            "rule_set_id",
            "reporting_currency",
        }
        present = run_config["config_key"].dropna().astype(str)
        for key in sorted(required_keys - set(present)):
            issues.append(
                ValidationIssue(
                    "ERROR",
                    "MISSING_RUN_CONFIG",
                    "run_config",
                    "",
                    "config_key",
                    f"Laufparameter fehlt: {key}",
                )
            )
        duplicates = present[present.duplicated(keep=False)]
        for key in sorted(set(duplicates)):
            issues.append(
                ValidationIssue(
                    "ERROR",
                    "DUPLICATE_RUN_CONFIG",
                    "run_config",
                    key,
                    "config_key",
                    "Laufparameter ist mehrdeutig",
                )
            )
    tranches = tables.get("securitisation_tranche")
    if tranches is not None and not tranches.empty:
        invalid = pd.to_numeric(tranches["attachment"], errors="coerce") >= pd.to_numeric(
            tranches["detachment"], errors="coerce"
        )
        for index in tranches.index[invalid][:20]:
            issues.append(
                ValidationIssue(
                    "ERROR",
                    "INVALID_TRANCHE_BOUNDS",
                    "securitisation_tranche",
                    str(tranches.at[index, "record_id"]),
                    "attachment,detachment",
                    "Es muss A < D gelten",
                )
            )
    sensitivities = tables.get("market_sensitivity")
    if sensitivities is not None and not sensitivities.empty:
        curvature = sensitivities["measure"].astype(str).eq("CURVATURE")
        missing = curvature & (sensitivities["curvature_up"].isna() | sensitivities["curvature_down"].isna())
        for index in sensitivities.index[missing][:20]:
            issues.append(
                ValidationIssue(
                    "ERROR",
                    "MISSING_CURVATURE_REVALUATION",
                    "market_sensitivity",
                    str(sensitivities.at[index, "record_id"]),
                    "curvature_up,curvature_down",
                    "Curvature verlangt Up- und Down-Revaluation",
                )
            )
    issues.extend(ValidationIssue(*issue) for issue in support_issues(tables))
    return issues


def _normalise_table_name(name: str) -> str:
    name = re.sub(r"[^A-Za-z0-9_]", "_", name)
    if not name or not name[0].isalpha():
        name = "t_" + name
    return name[:250]


def write_result_workbook(
    path: Path, sheets: Mapping[str, pd.DataFrame], *, metadata: Mapping[str, object]
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(
        path, engine="xlsxwriter", datetime_format="yyyy-mm-dd", date_format="yyyy-mm-dd"
    ) as writer:
        meta = pd.DataFrame([{"key": k, "value": v} for k, v in metadata.items()])
        meta.to_excel(writer, sheet_name="Run_Info", index=False)
        workbook = writer.book
        header = workbook.add_format(
            {"bold": True, "font_color": "white", "bg_color": "#1F4E78", "border": 1, "align": "center"}
        )
        money = workbook.add_format({"num_format": "#,##0.00;[Red]-#,##0.00"})
        rate = workbook.add_format({"num_format": "0.0000%"})
        date_fmt = workbook.add_format({"num_format": "yyyy-mm-dd"})
        info_ws = writer.sheets["Run_Info"]
        info_ws.freeze_panes(1, 0)
        info_ws.add_table(
            0,
            0,
            max(len(meta), 1),
            1,
            {
                "name": "tbl_out_run_info",
                "style": "Table Style Medium 2",
                "columns": [{"header": "key"}, {"header": "value"}],
            },
        )
        info_ws.set_column(0, 0, 28)
        info_ws.set_column(1, 1, 70)
        info_ws.set_row(0, None, header)
        for sheet_name, frame in sheets.items():
            safe_sheet = sheet_name[:31]
            clean = frame.copy()
            for col in clean.columns:
                if clean[col].dtype == object:
                    clean[col] = clean[col].map(_safe_excel_value)
            clean.to_excel(writer, sheet_name=safe_sheet, index=False)
            ws = writer.sheets[safe_sheet]
            ws.freeze_panes(1, 0)
            nrows, ncols = clean.shape
            if ncols:
                ws.add_table(
                    0,
                    0,
                    max(nrows, 1),
                    ncols - 1,
                    {
                        "name": _normalise_table_name(f"tbl_out_{safe_sheet}"),
                        "style": "Table Style Medium 2",
                        "columns": [{"header": c} for c in clean.columns],
                    },
                )
            for idx, col in enumerate(clean.columns):
                width = min(
                    max(
                        len(str(col)) + 2,
                        max((len(str(v)) for v in clean[col].head(200).dropna()), default=0) + 2,
                    ),
                    40,
                )
                fmt = None
                lower = str(col).lower()
                if any(
                    x in lower for x in ("rate", "ratio", "share", "factor", "weight", "pd", "lgd", "ccf")
                ):
                    fmt = rate
                elif any(
                    x in lower
                    for x in ("amount", "ead", "rwea", "trea", "capital", "eve", "nii", "value", "loss")
                ):
                    fmt = money
                elif "date" in lower or lower.endswith("_at"):
                    fmt = date_fmt
                ws.set_column(idx, idx, width, fmt)
            ws.set_row(0, None, header)
